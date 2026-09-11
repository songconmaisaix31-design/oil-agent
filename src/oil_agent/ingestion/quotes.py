"""Nonexecuting CSV/XLSX preview. Persistence and confirmation belong to C."""

import csv
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath

from openpyxl import load_workbook
from pydantic import ValidationError

from oil_agent.contracts.dto import (
    EvidenceRef,
    MarketObservation,
    Provenance,
    QualityState,
    SourceRecord,
    TimeQuality,
)
from oil_agent.contracts.services import ErrorCode, ServiceError
from oil_agent.ingestion.common import canonical_json, content_hash, stable_id

COMPARISON_FIELDS = (
    "product",
    "spec",
    "region",
    "supplier",
    "quote_type",
    "tax_basis",
    "delivery_basis",
    "currency",
    "unit",
)
FIELDS = (*COMPARISON_FIELDS, "value", "as_of", "published_at")


def comparison_key(observation: MarketObservation) -> tuple[str, ...] | None:
    values = tuple(getattr(observation, name) for name in COMPARISON_FIELDS)
    if any(
        v is None or not v.strip() or v.casefold() in {"unknown", "n/a", "未知"} for v in values
    ):
        return None
    if observation.period_start or observation.quality_state in {
        QualityState.INVALID,
        QualityState.INCOMPLETE,
    }:
        return None
    return values


@dataclass(frozen=True)
class UploadLimits:
    max_bytes: int = 2_000_000
    max_expanded_bytes: int = 10_000_000
    max_entries: int = 200
    max_ratio: int = 100
    max_rows: int = 10000
    max_columns: int = 64
    max_cell_chars: int = 2000

    def __post_init__(self):
        if any(not isinstance(v, int) or not 1 <= v <= 20_000_000 for v in vars(self).values()):
            raise ValueError("Upload limits must be finite positive integers")


@dataclass(frozen=True)
class PreviewRow:
    row_number: int
    row_id: str
    mapped: dict[str, str]
    errors: tuple[str, ...]
    duplicate_of: int | None = None
    record: SourceRecord | None = None
    observation: MarketObservation | None = None


@dataclass(frozen=True)
class QuotePreview:
    file_id: str
    file_sha256: str
    columns: tuple[str, ...]
    rows: tuple[PreviewRow, ...]


def _reject(message: str):
    raise ServiceError(ErrorCode.INVALID_INPUT, message)


def _xlsx_rows(data: bytes, limits: UploadLimits) -> list[list[str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > limits.max_entries or len({i.filename for i in infos}) != len(infos):
                _reject("XLSX has excessive or duplicate archive entries")
            if sum(i.file_size for i in infos) > limits.max_expanded_bytes:
                _reject("XLSX expansion exceeds limit")
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= {i.filename for i in infos}:
                _reject("Upload is not an XLSX workbook")
            for info in infos:
                path = PurePosixPath(info.filename)
                name = info.filename.casefold()
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or "\\" in name
                    or ":" in name
                    or info.flag_bits & 1
                    or info.compress_type not in (0, 8)
                    or info.file_size / max(1, info.compress_size) > limits.max_ratio
                    or any(
                        t in name
                        for t in (
                            "vbaproject",
                            "macrosheet",
                            "embeddings/",
                            "externallinks/",
                            "activex",
                        )
                    )
                ):
                    _reject("Unsafe XLSX archive entry")
                if name.endswith((".xml", ".rels")):
                    payload = archive.read(info).lower()
                    if any(
                        t in payload
                        for t in (
                            b"<!doctype",
                            b"<!entity",
                            b"macroenabled",
                            b' targetmode="external"',
                            b" targetmode='external'",
                        )
                    ):
                        _reject("External, macro or entity content is forbidden")
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(workbook.worksheets) != 1:
                _reject("Preview requires one worksheet; select a sheet before upload")
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            rows = []
            for cells in sheet.iter_rows():
                if len(rows) > limits.max_rows or len(cells) > limits.max_columns:
                    _reject("Worksheet row/column limit exceeded")
                rows.append(
                    [
                        "=" + str(c.value)
                        if c.data_type == "f" and not str(c.value).startswith("=")
                        else (
                            c.value.isoformat() if isinstance(c.value, datetime) else str(c.value)
                        )
                        if c.value is not None
                        else ""
                        for c in cells
                    ]
                )
            return rows
        finally:
            workbook.close()
    except ServiceError:
        raise
    except Exception:
        _reject("Invalid or unsupported XLSX workbook")


def _rows(data: bytes, filename: str, limits: UploadLimits) -> list[list[str]]:
    if not data or len(data) > limits.max_bytes:
        _reject("Upload is empty or exceeds size limit")
    extension = PurePosixPath(filename).suffix.casefold()
    if extension == ".xlsx":
        if not data.startswith(b"PK\x03\x04"):
            _reject("XLSX extension does not match content")
        return _xlsx_rows(data, limits)
    if extension != ".csv" or data.startswith((b"PK", b"\xd0\xcf\x11\xe0")) or b"\x00" in data:
        _reject("Only UTF-8 CSV and non-macro XLSX uploads are supported")
    try:
        parsed = []
        for row in csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True):
            if len(parsed) > limits.max_rows or len(row) > limits.max_columns:
                _reject("CSV row/column limit exceeded")
            parsed.append(row)
        return parsed
    except (UnicodeError, csv.Error):
        _reject("CSV must be valid UTF-8 with valid quoting")


def _time(raw: str) -> datetime:
    value = datetime.fromisoformat(raw)
    if value.tzinfo is None:
        raise ValueError("Timestamp requires timezone")
    return value.astimezone(UTC)


def preview_quotes(
    data: bytes,
    filename: str,
    mapping: dict[str, str],
    *,
    rights_ref: str,
    origin_publisher: str,
    discovered_at: datetime,
    is_fixture: bool,
    provenance: Provenance,
    fixture_dataset: str | None = None,
    limits: UploadLimits | None = None,
) -> QuotePreview:
    """Mapping is canonical field -> column heading; missing basis stays uncomparable.

    An empty mapping selects exact canonical names present in the validated header.
    File identity hashes original bytes; row identity includes file, row index and
    resolved mapping. Identical explicit/derived mappings produce identical IDs.
    """
    limits = limits or UploadLimits()
    if discovered_at.tzinfo is None or not rights_ref.strip() or not origin_publisher.strip():
        _reject("Timezone, rights reference and publisher are required")
    raw_rows = _rows(data, filename, limits)
    if not raw_rows:
        _reject("Upload needs a header")
    header = tuple(c.strip() for c in raw_rows[0])
    if any(c.startswith(("=", "+", "-", "@")) or len(c) > limits.max_cell_chars for c in header):
        _reject("Unsafe or oversized column heading")
    if len(set(header)) != len(header) or any(not c for c in header):
        _reject("Column headings must be unique and nonempty")
    if not mapping:
        mapping = {field: field for field in FIELDS if field in header}
    if set(mapping) - set(FIELDS) or not {"value", "as_of"} <= set(mapping):
        _reject("Mapping requires value and as_of and must use known fields")
    if len(set(mapping.values())) != len(mapping):
        _reject("A column cannot map to several fields")
    if set(mapping.values()) - set(header):
        _reject("Mapped column does not exist")
    digest = hashlib.sha256(data).hexdigest()
    rows, seen = [], {}
    for index, cells in enumerate(raw_rows[1:], 2):
        row_id = stable_id("quote-row", digest, index, mapping)
        mapped = {
            field: cells[header.index(column)].strip() if header.index(column) < len(cells) else ""
            for field, column in mapping.items()
        }
        errors = []
        if len(cells) != len(header):
            errors.append("column_count_mismatch")
        if any(len(c) > limits.max_cell_chars for c in cells):
            errors.append("cell_too_large")
        if any(c.lstrip().startswith(("=", "+", "-", "@")) for c in cells):
            errors.append("formula_or_executable_cell")
        key = canonical_json(mapped)
        duplicate = seen.get(key)
        if duplicate is not None:
            errors.append("duplicate_row")
        else:
            seen[key] = index
        record = observation = None
        if not errors:
            try:
                value = Decimal(mapped["value"])
                at = _time(mapped["as_of"])
                published = _time(mapped["published_at"]) if mapped.get("published_at") else None
                title = "Uploaded quote row"
                payload = {k: mapped.get(k) or None for k in FIELDS}
                payload.update(value=str(value), as_of=at.isoformat())
                excerpt = canonical_json(payload)
                if len(excerpt) > 2000:
                    raise ValueError("Row exceeds evidence size bound")
                fixture = dict(
                    is_fixture=is_fixture, provenance=provenance, fixture_dataset=fixture_dataset
                )
                future = at > discovered_at or (published is not None and published > discovered_at)
                record = SourceRecord(
                    record_id=row_id,
                    source_id="quote-upload",
                    external_id=row_id,
                    revision=1,
                    title=title,
                    content_excerpt=excerpt,
                    content_hash=content_hash(title, excerpt),
                    url=None,
                    origin_publisher=origin_publisher,
                    rights_ref=rights_ref,
                    published_at=published,
                    discovered_at=discovered_at,
                    time_quality=TimeQuality.FUTURE_QUARANTINED if future else TimeQuality.VALID,
                    **fixture,
                )
                observation = MarketObservation(
                    observation_id=stable_id("quote", row_id),
                    revision=1,
                    **{
                        k: mapped.get(k) or None
                        for k in ("product", "spec", "region", "supplier", "currency", "unit")
                    },
                    quote_type=mapped.get("quote_type") or "unknown",
                    tax_basis=mapped.get("tax_basis") or "unknown",
                    delivery_basis=mapped.get("delivery_basis") or "unknown",
                    value=value,
                    as_of=at,
                    published_at=published,
                    source_record_id=row_id,
                    evidence=EvidenceRef(
                        record_id=row_id, revision=1, field="content_excerpt", excerpt=excerpt
                    ),
                    quality_state=QualityState.INVALID if future else QualityState.VALID,
                    **fixture,
                )
                if not future and comparison_key(observation) is None:
                    observation = observation.model_copy(
                        update={"quality_state": QualityState.INCOMPLETE}
                    )
            except (InvalidOperation, ValueError, ValidationError):
                errors.append("invalid_value_time_or_basis")
                record = observation = None
        rows.append(
            PreviewRow(index, row_id, mapped, tuple(errors), duplicate, record, observation)
        )
    return QuotePreview("quote-file:" + digest, digest, header, tuple(rows))


def validate_observation(observation: MarketObservation, record: SourceRecord) -> bool:
    """Verify that numeric and basis fields equal the exact structured source row."""
    ref = observation.evidence
    if (ref.record_id, ref.revision, ref.field) != (
        record.record_id,
        record.revision,
        "content_excerpt",
    ):
        return False
    if ref.excerpt != record.content_excerpt or observation.is_fixture != record.is_fixture:
        return False
    try:
        payload = json.loads(ref.excerpt)
        fields = (
            *COMPARISON_FIELDS,
            "value",
            "as_of",
            "published_at",
            "period_start",
            "period_end",
        )
        for field in fields:
            actual, raw = getattr(observation, field), payload.get(field)
            if field == "value":
                if isinstance(raw, (float, bool)) or Decimal(raw) != actual:
                    return False
            elif field in ("as_of", "published_at") and raw:
                if _time(raw) != actual:
                    return False
            elif field in ("period_start", "period_end") and actual:
                if actual.isoformat() != raw:
                    return False
            elif actual in ("unknown", None) and raw is None:
                continue
            elif actual != raw:
                return False
        return True
    except (ValueError, TypeError, KeyError, InvalidOperation):
        return False
